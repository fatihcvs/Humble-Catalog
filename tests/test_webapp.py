import io
import json
import re
from pathlib import Path
import requests
from unittest.mock import Mock
from openpyxl import load_workbook
from humble_catalog import db, export, stats
from humble_catalog.webapp import create_app


def _viewer_js():
    """Every viewer script concatenated, in load order.

    These assertions pin that the viewer does something, not that one
    file does. Reading app.js alone made them break when a function moved
    between scripts, which is a fact about the file layout and not about
    the behaviour they were written to protect.
    """
    from tests.js_harness import VIEWER_JS
    return "\n".join(p.read_text(encoding="utf-8") for p in VIEWER_JS)


def test_index_offers_android_type_filter():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert '<option value="android">Android apps</option>' in html

def test_index_has_autocomplete_filters():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert '<input id="f-genre"' in html      # selects replaced by
    assert '<input id="f-bundle"' in html     # autocomplete inputs
    assert '<select id="f-type">' in html     # small closed lists stay selects
    assert '<select id="f-flag">' in html

def test_index_offers_the_annotation_flags():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    # the values are what app.js switches on; the labels are what is read
    assert '<option value="notes">Has notes</option>' in html
    assert '<option value="mytags">Has my tags</option>' in html

def test_index_has_person_filters():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    # authors and narrator are array fields, so they get chip filters
    # like genre/bundle rather than living only in the search haystack
    assert '<input id="f-authors"' in html
    assert '<input id="f-narrator"' in html
    assert '<input id="f-publisher"' in html

def test_index_marks_narrator_and_bundle_sortable():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert 'data-sort="narrator"' in html
    assert 'data-sort="bundle"' in html
    assert 'data-sort="read_status"' in html
    # every sortable header carries an indicator slot (Status added an 11th)
    assert html.count('class="sort-ind"') == 11

def test_only_the_table_scrolls_sideways():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    html = (static / "index.html").read_text(encoding="utf-8")
    css = (static / "style.css").read_text(encoding="utf-8")
    # the 13 columns outrun the viewport, so something has to scroll. A
    # scroller around <main> was tried and dragged the panels sideways
    # with the table, which reads as the whole page scrolling; the
    # scroller has to wrap the table alone.
    assert '<div id="table-wrap">' in html
    assert "#table-wrap { flex: 1; min-height: 0; overflow: auto; }" in css
    assert "body { margin: 0" in css and "overflow: hidden;" in css
    # #table-wrap scrolls vertically too, so the sticky <thead> has a
    # scrollport of its own to stick within
    assert "#catalog thead th { position: sticky; top: 0;" in css
    # panels sit outside that scroller, so they need a cap of their own or
    # an expanded one squeezes the table region to nothing
    # the cap is gone with the stacking that needed it: a section owns the
    # viewport, so no panel can squeeze the table region any more
    assert "max-height: 50%" not in css
    assert 'section[id^="section-"]' in css

def test_app_wires_every_registered_chip_filter():
    js = _viewer_js()
    for field in ("genre", "series", "authors", "narrator", "publisher",
                  "bundle"):
        assert f"{field}:" in js
    # wiring loops over the registry, so an entry cannot end up as a
    # silently dead control the way a hand-written call list allowed
    assert "for (const field of Object.keys(chipFilters)) wireChipFilter(field);" in js

def test_narrator_filter_spans_illustrator():
    js = _viewer_js()
    # the Narrator/Artist column shows narrator || illustrator, so a
    # narrator-only filter would silently miss every comic illustrator
    assert "[...i.narrator, ...i.illustrator]" in js

def test_search_matches_name_only():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    js = _viewer_js()
    # every other field the old haystack spanned now has its own filter.
    # Matching became fuzzy (fuzzy.js), but the haystack is still the
    # name and nothing else.
    assert "Fuzzy.score(q, i.name" in js
    assert "...i.bundles.map(b => b.name)].join" not in js
    html = (static / "index.html").read_text(encoding="utf-8")
    # the word "Search" is a standalone label, not placeholder text
    assert '<label id="search-label" for="search">Search</label>' in html
    assert 'placeholder="Name..."' in html

def test_series_filter_suppresses_all_any_toggle():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    html = (static / "index.html").read_text(encoding="utf-8")
    assert '<input id="f-series"' in html
    js = _viewer_js()
    # series is one-per-item, so "all" with 2+ chips is unsatisfiable:
    # the flag hides a toggle that could only ever empty the table
    assert "scalar: true" in js
    assert "!f.scalar" in js

def test_autocomplete_opens_only_on_typing_or_arrows():
    ac = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
          / "static" / "autocomplete.js").read_text(encoding="utf-8")
    # opening on focus made the list flash up whenever a field was
    # clicked; the arrows summon it deliberately instead
    assert 'addEventListener("focus"' not in ac
    assert "if (!owned()) show();" in ac

def test_autocomplete_list_belongs_to_one_input():
    ac = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
          / "static" / "autocomplete.js").read_text(encoding="utf-8")
    # one list is shared across every attached input, so each handler
    # must check the open one is its own: otherwise a blur timer closes
    # a list another input just opened, and arrow keys drive a list
    # whose values belong to a different field
    assert "openOwner" in ac
    assert "const owned = () => openList && openOwner === input;" in ac
    assert "if (owned()) close();" in ac

def test_autocomplete_popup_is_not_a_tab_stop():
    ac = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
          / "static" / "autocomplete.js").read_text(encoding="utf-8")
    # the list scrolls and has no focusable children, so Chrome would
    # otherwise hand it a tab stop of its own as a "focusable scroller",
    # stealing the Tab that should reach the next filter field
    assert 'openList.tabIndex = -1;' in ac

def test_autocomplete_anchors_popup_to_the_input():
    ac = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
          / "static" / "autocomplete.js").read_text(encoding="utf-8")
    # .ac-wrap also holds chips, which wrap onto extra lines; without an
    # explicit anchor the list renders at its static position and drifts
    assert "openOwner.getBoundingClientRect()" in ac
    assert "openList.style.top" in ac and "openList.style.left" in ac
    # the list hangs off <body> so <main>'s overflow cannot clip it, which
    # means nothing moves it with its anchor unless we do it ourselves
    assert "document.body.appendChild(openList)" in ac
    assert 'window.addEventListener("scroll", reposition, true)' in ac
    assert 'window.addEventListener("resize", reposition)' in ac
    # an in-cell list outlives the row rebuild that used to remove it
    assert "if (!openOwner.isConnected) { close(); return; }" in ac
    css = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
           / "static" / "style.css").read_text(encoding="utf-8")
    # the list may grow rightwards past the field's own width
    assert "width: max-content" in css
    # viewport coordinates only work against a fixed element
    assert ".ac-list { position: fixed;" in css

def test_search_box_has_title_typeahead():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    html = (static / "index.html").read_text(encoding="utf-8")
    # the dropdown positions against a .ac-wrap, so #search must sit
    # inside one; the wrapper takes over the flex sizing
    assert 'id="search-wrap"' in html
    css = (static / "style.css").read_text(encoding="utf-8")
    assert "#search-wrap" in css
    js = _viewer_js()
    # one stray keystroke should not open a list drawn from every title
    assert "length < 2" in js
    assert 'Autocomplete.attach(\n  $("#search")' in js

def _seed(dbp):
    conn = db.connect(dbp)
    conn.execute("INSERT INTO bundles VALUES ('k1','Bundle One','http://b1','2020-01-01')")
    cur = conn.execute(
        "INSERT INTO items (machine_name, name, type, publisher) "
        "VALUES ('asr','All Systems Red','ebook','Example Press')")
    item_id = cur.lastrowid
    conn.execute("INSERT INTO item_bundles VALUES (?, 'k1')", (item_id,))
    cands = [{"source": "hardcover", "title": "All Systems Red",
              "authors": ["Martha Wells"], "genre": "SF", "series": None,
              "series_number": None, "rating": 4.3, "narrator": None,
              "illustrator": None, "extra": {}, "confidence": 0.7}]
    conn.execute(
        "INSERT INTO enrichment (item_id, status, candidates) VALUES (?,?,?)",
        (item_id, "low_confidence", json.dumps(cands)))
    conn.commit()
    conn.close()
    return item_id

def test_items_rating_and_review_flow(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    items = client.get("/api/items").get_json()["items"]
    assert items[0]["name"] == "All Systems Red"
    assert items[0]["bundles"][0]["name"] == "Bundle One"

    assert client.post(f"/api/items/{item_id}/rating",
                       json={"rating": 5}).status_code == 200
    assert client.get("/api/items").get_json()["items"][0]["my_rating"] == 5

    review = client.get("/api/review").get_json()["items"]
    assert review and review[0]["candidates"][0]["title"] == "All Systems Red"

    assert client.post(f"/api/items/{item_id}/choose",
                       json={"candidate": 0}).status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["status"] == "manually_fixed" and item["genre"] == ["SF"]
    assert client.get("/api/review").get_json()["items"] == []


def test_set_read_status_updates_the_item(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post(f"/api/items/{item_id}/read-status", json={"status": "reading"})
    assert resp.status_code == 200
    conn = db.connect(dbp)
    assert conn.execute("SELECT read_status FROM items WHERE id=?",
                        (item_id,)).fetchone()["read_status"] == "reading"
    conn.close()


def test_set_read_status_rejects_an_unknown_value(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post(f"/api/items/{item_id}/read-status",
                       json={"status": "skimmed"}).status_code == 400


def test_set_read_status_404s_for_a_missing_item(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/items/99999/read-status",
                       json={"status": "read"}).status_code == 404

def test_type_override_and_status(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post(f"/api/items/{item_id}/type",
                       json={"type": "comic"}).status_code == 200
    assert client.get("/api/items").get_json()["items"][0]["type"] == "comic"
    assert client.get("/api/status").get_json() == {"runs": []}

def test_review_sorted_by_confidence_with_covers(tmp_path):
    dbp = tmp_path / "t.db"
    conn = db.connect(dbp)
    conn.execute("INSERT INTO items (machine_name, name, type, cover_path) "
                 "VALUES ('a','Low Conf Book','ebook','covers/1.jpg')")
    conn.execute("INSERT INTO enrichment (item_id, status, candidates) VALUES "
                 "(1,'low_confidence','[{\"source\":\"s\",\"title\":\"A\",\"confidence\":0.62}]')")
    conn.execute("INSERT INTO items (machine_name, name, type) VALUES ('b','High Conf Book','ebook')")
    conn.execute("INSERT INTO enrichment (item_id, status, candidates) VALUES "
                 "(2,'low_confidence','[{\"source\":\"s\",\"title\":\"B1\",\"confidence\":0.4},"
                 "{\"source\":\"s\",\"title\":\"B2\",\"confidence\":0.81}]')")
    conn.commit()
    conn.close()
    client = create_app(db_path=str(dbp)).test_client()
    review = client.get("/api/review").get_json()["items"]
    assert [r["name"] for r in review] == ["High Conf Book", "Low Conf Book"]
    assert review[0]["candidates"][0]["title"] == "B2"  # sorted within item too
    assert review[1]["cover_path"] == "covers/1.jpg"

def test_choose_index_matches_review_order(tmp_path):
    # Candidates stored low-confidence-first; /api/review shows them sorted
    # desc, and the frontend posts an index into that sorted view. Choosing
    # index 0 must apply the highest-confidence candidate, not stored[0].
    dbp = tmp_path / "t.db"
    conn = db.connect(dbp)
    conn.execute("INSERT INTO items (machine_name, name, type) VALUES ('x','X','ebook')")
    cands = [{"source": "s", "title": "Woodworking X", "genre": "Crafts",
              "confidence": 0.55},
             {"source": "s", "title": "Programming X", "genre": "Computing",
              "confidence": 0.84}]
    conn.execute("INSERT INTO enrichment (item_id, status, candidates) "
                 "VALUES (1,'low_confidence',?)", (json.dumps(cands),))
    conn.commit()
    conn.close()
    client = create_app(db_path=str(dbp)).test_client()
    review = client.get("/api/review").get_json()["items"]
    assert review[0]["candidates"][0]["title"] == "Programming X"
    assert client.post("/api/items/1/choose",
                       json={"candidate": 0}).status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["genre"] == ["Computing"]

def test_reopen_and_apply_endpoints(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/choose", json={"candidate": 0})
    assert client.get("/api/review").get_json()["items"] == []

    # reopen: back in the queue, data and candidates intact
    assert client.post(f"/api/items/{item_id}/reopen").status_code == 200
    review = client.get("/api/review").get_json()["items"]
    assert review and review[0]["candidates"]

    # apply an arbitrary candidate (as the URL importer will)
    cand = {"source": "comicvine", "title": "Shadow Hound: Origins",
            "genre": "Manga", "url": "https://comicvine.gamespot.com/x/4050-1/",
            "authors": ["Bo Writer"]}
    assert client.post(f"/api/items/{item_id}/apply",
                       json={"candidate": cand}).status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["status"] == "manually_fixed"
    assert item["genre"] == ["Manga"]
    assert item["source_url"] == "https://comicvine.gamespot.com/x/4050-1/"
    assert client.post(f"/api/items/{item_id}/apply",
                       json={"candidate": {"title": "no source"}}).status_code == 400

def test_fetch_url_endpoint(tmp_path, monkeypatch):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    assert client.post(f"/api/items/{item_id}/fetch_url",
                       json={}).status_code == 400

    from humble_catalog import url_import
    fake = {"source": "drivethrurpg", "title": "Heart", "url": "https://d/x"}
    monkeypatch.setattr(url_import, "resolve", lambda conn, url: fake)
    resp = client.post(f"/api/items/{item_id}/fetch_url",
                       json={"url": "https://www.drivethrurpg.com/en/product/1/x"})
    assert resp.status_code == 200
    assert resp.get_json()["candidate"]["title"] == "Heart"

    def boom(conn, url):
        raise ValueError("unsupported source URL")
    monkeypatch.setattr(url_import, "resolve", boom)
    resp = client.post(f"/api/items/{item_id}/fetch_url",
                       json={"url": "https://example.com/x"})
    assert resp.status_code == 400
    assert "unsupported" in resp.get_json()["error"]

def test_index_has_the_stats_panel():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    html = (static / "index.html").read_text(encoding="utf-8")
    assert '<div id="stats-panel" hidden></div>' in html
    # the two panels it replaced are gone, not merely hidden
    assert "genres-panel" not in html and "gaps-panel" not in html
    # genre management survives inside it, behind the Edit tags toggle
    js = _viewer_js()
    assert "/api/genres/rename" in js and "/api/genres/delete" in js

def _seed_genres_app(tmp_path):
    dbp = tmp_path / "g.db"
    conn = db.connect(dbp)
    for idx, (mn, name) in enumerate(
            [("wad", "Wings of Autumn Dusk"), ("ub", "Unrelated Book")], 1):
        conn.execute("INSERT INTO items (machine_name, name) VALUES (?,?)",
                     (mn, name))
        conn.execute("INSERT INTO enrichment (item_id, genre) VALUES (?,?)",
                     (idx, '["Fantasy"]' if idx == 1 else '["Horror"]'))
    conn.commit()
    conn.close()
    return create_app(db_path=str(dbp)).test_client()

def test_genre_rename_endpoint(tmp_path):
    client = _seed_genres_app(tmp_path)
    resp = client.post("/api/genres/rename",
                       json={"old": "Fantasy", "new": "Epic Fantasy"})
    assert resp.status_code == 200 and resp.get_json() == {"changed": 1}
    genres = [i["genre"] for i in client.get("/api/items").get_json()["items"]]
    assert ["Epic Fantasy"] in genres and ["Fantasy"] not in genres

def test_genre_rename_endpoint_rejects_bad_input(tmp_path):
    client = _seed_genres_app(tmp_path)
    assert client.post("/api/genres/rename",
                       json={"old": "", "new": "X"}).status_code == 400
    assert client.post("/api/genres/rename",
                       json={"old": "Fantasy"}).status_code == 400
    assert client.post("/api/genres/rename",
                       json={"old": "Cooking", "new": "Food"}).status_code == 404

def test_genre_delete_endpoint(tmp_path):
    client = _seed_genres_app(tmp_path)
    resp = client.post("/api/genres/delete", json={"tag": "Horror"})
    assert resp.status_code == 200 and resp.get_json() == {"changed": 1}
    genres = [i["genre"] for i in client.get("/api/items").get_json()["items"]]
    assert [] in genres and ["Horror"] not in genres

def test_genre_delete_endpoint_rejects_bad_input(tmp_path):
    client = _seed_genres_app(tmp_path)
    assert client.post("/api/genres/delete", json={}).status_code == 400
    assert client.post("/api/genres/delete",
                       json={"tag": "Cooking"}).status_code == 404

def test_edit_snapshots_once_and_updates(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/choose", json={"candidate": 0})

    # first edit: snapshot taken, fields updated, status untouched
    assert client.post(f"/api/items/{item_id}/edit", json={"fields": {
        "genre": ["Cyberpunk", "Noir"], "series_number": "2"}}).status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["genre"] == ["Cyberpunk", "Noir"] and item["series_number"] == 2.0
    assert item["edited"] is True
    assert item["status"] == "manually_fixed"

    # second edit must NOT re-snapshot (revert target stays the enriched form)
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"genre": ["Solarpunk"], "authors": []}})
    item = client.get("/api/items").get_json()["items"][0]
    assert item["genre"] == ["Solarpunk"] and item["authors"] == []

    assert client.post(f"/api/items/{item_id}/revert").status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["genre"] == ["SF"]              # back to the enriched values
    assert item["authors"] == ["Martha Wells"]
    assert item["edited"] is False

def test_edit_validation(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post(f"/api/items/{item_id}/edit",
                       json={"fields": {"status": "matched"}})
    assert resp.status_code == 400 and "status" in resp.get_json()["error"]
    assert client.post(f"/api/items/{item_id}/edit", json={"fields": {
        "series_number": "two"}}).status_code == 400
    assert client.post(f"/api/items/{item_id}/edit",
                       json={}).status_code == 400
    assert client.post(f"/api/items/{item_id}/edit", json={"fields": {
        "genre": "not a list"}}).status_code == 400
    assert client.post(f"/api/items/{item_id}/edit", json={"fields": {
        "authors": ["", "  "]}}).status_code == 200  # all-blank list stores NULL

def test_revert_without_edit_is_400(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post(f"/api/items/{item_id}/revert").status_code == 400

def test_apply_candidate_keeps_a_revert_target_for_a_hand_edit(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"genre": ["Hand Typed"]}})
    assert client.get("/api/items").get_json()["items"][0]["edited"] is True
    # Applying a candidate hands authorship back to enrichment, but the
    # typed values become the new revert target instead of being dropped.
    client.post(f"/api/items/{item_id}/choose", json={"candidate": 0})
    item = client.get("/api/items").get_json()["items"][0]
    assert item["edited"] is False and item["genre"] == ["SF"]
    assert item["re_enriched"] is True
    assert client.post(f"/api/items/{item_id}/revert").status_code == 200
    assert client.get("/api/items").get_json()["items"][0]["genre"] == ["Hand Typed"]

def test_index_serves_html(tmp_path):
    _seed(tmp_path / "t.db")
    client = create_app(db_path=str(tmp_path / "t.db")).test_client()
    resp = client.get("/")
    assert resp.status_code == 200
    assert b"<table" in resp.data or b"catalog" in resp.data.lower()

def test_export_csv_endpoint(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.get("/api/export.csv")
    assert resp.status_code == 200
    assert resp.mimetype == "text/csv"
    assert resp.headers["Content-Disposition"] == "attachment; filename=catalog.csv"
    body = resp.get_data()
    assert body.startswith(b"\xef\xbb\xbftitle,")  # UTF-8 BOM so Excel detects encoding
    assert b"All Systems Red" in body

# The POST tests below use _seed_pair, defined just after them: the export
# tests stay together rather than splitting around a shared helper.

def test_export_csv_post_uses_the_posted_ids_in_order(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.csv", json={"ids": list(reversed(ids))})
    assert resp.status_code == 200
    assert resp.mimetype == "text/csv"
    assert resp.headers["Content-Disposition"] == "attachment; filename=catalog.csv"
    body = resp.get_data()
    assert body.startswith(b"\xef\xbb\xbftitle,")  # same BOM as the GET path
    rows = body.decode("utf-8-sig").splitlines()
    assert len(rows) == 3  # header + 2 items
    # the second seeded item's row comes first, because that is what was posted
    first_data_row = rows[1]
    resp_fwd = client.post("/api/export.csv", json={"ids": ids})
    fwd_rows = resp_fwd.get_data().decode("utf-8-sig").splitlines()
    assert fwd_rows[1] != first_data_row

def test_export_csv_post_subset(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.csv", json={"ids": [ids[0]]})
    rows = resp.get_data().decode("utf-8-sig").splitlines()
    assert len(rows) == 2  # header + 1 item

def test_export_csv_post_rejects_a_bad_body(tmp_path):
    dbp = tmp_path / "t.db"
    _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/export.csv", json={}).status_code == 400
    assert client.post("/api/export.csv", json={"ids": "all"}).status_code == 400

def test_export_csv_get_still_returns_the_whole_catalog(tmp_path):
    # GET keeps meaning "everything": bookmarks, curl, and scripts use it
    # even though the viewer now posts.
    dbp = tmp_path / "t.db"
    _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    rows = client.get("/api/export.csv").get_data().decode("utf-8-sig").splitlines()
    assert len(rows) == 3  # header + both items

_XLSX_MIME = ("application/vnd.openxmlformats-officedocument"
              ".spreadsheetml.sheet")

def _xlsx_titles(resp):
    """Data-row titles from an xlsx response, in file order."""
    ws = load_workbook(io.BytesIO(resp.get_data())).active
    col = [c.value for c in ws[1]].index("title")
    return [row[col].value for row in ws.iter_rows(min_row=2)]

def test_export_xlsx_post_uses_the_posted_ids_in_order(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.xlsx", json={"ids": list(reversed(ids))})
    assert resp.status_code == 200
    assert resp.mimetype == _XLSX_MIME
    assert resp.headers["Content-Disposition"] == \
        "attachment; filename=catalog.xlsx"
    forward = _xlsx_titles(client.post("/api/export.xlsx", json={"ids": ids}))
    assert _xlsx_titles(resp) == list(reversed(forward))

def test_export_xlsx_get_returns_the_whole_catalog(tmp_path):
    # GET keeps meaning "everything", same as the CSV route: it is the
    # form a bookmark or a script can use.
    dbp = tmp_path / "t.db"
    _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert len(_xlsx_titles(client.get("/api/export.xlsx"))) == 2

def test_export_xlsx_post_rejects_a_bad_body(tmp_path):
    dbp = tmp_path / "t.db"
    _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/export.xlsx", json={}).status_code == 400
    assert client.post("/api/export.xlsx", json={"ids": "all"}).status_code == 400

def _seed_pair(dbp, type_b="ebook"):
    conn = db.connect(dbp)
    conn.execute("INSERT INTO bundles VALUES ('k1','Bundle One','http://b1',NULL)")
    conn.execute("INSERT INTO bundles VALUES ('k2','Bundle Two','http://b2',NULL)")
    ids = []
    for mn, name, typ, gk in [("m1", "Book 2e", "ebook", "k1"),
                              ("m2", "Book, 2nd Edition", type_b, "k2")]:
        cur = conn.execute(
            "INSERT INTO items (machine_name, name, type) VALUES (?,?,?)",
            (mn, name, typ))
        conn.execute("INSERT INTO enrichment (item_id, genre) VALUES (?,?)",
                     (cur.lastrowid, db.tags_to_json(["Tech"])))
        conn.execute("INSERT INTO item_bundles VALUES (?,?)", (cur.lastrowid, gk))
        ids.append(cur.lastrowid)
    conn.commit()
    conn.close()
    return ids

def test_duplicates_endpoint_groups_and_dismissal(tmp_path):
    dbp = tmp_path / "t.db"
    a, b = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    groups = client.get("/api/duplicates").get_json()["groups"]
    assert len(groups) == 1
    member = groups[0][0]
    assert member["id"] == a and member["genre"] == ["Tech"]
    assert member["bundles"] == ["Bundle One"] and member["edited"] is False
    assert client.post("/api/dismiss_pair",
                       json={"id_a": a, "id_b": b}).status_code == 200
    assert client.get("/api/duplicates").get_json()["groups"] == []

def test_merge_endpoint_and_guards(tmp_path):
    dbp = tmp_path / "t.db"
    a, b = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/merge", json={"keep_id": a, "drop_id": a}).status_code == 400
    assert client.post("/api/merge", json={"keep_id": a, "drop_id": 999}).status_code == 400
    assert client.post("/api/merge", json={"keep_id": a, "drop_id": b}).status_code == 200
    items = client.get("/api/items").get_json()["items"]
    assert len(items) == 1
    assert {bu["name"] for bu in items[0]["bundles"]} == {"Bundle One", "Bundle Two"}

def test_merge_endpoint_rejects_cross_type(tmp_path):
    dbp = tmp_path / "t.db"
    a, b = _seed_pair(dbp, type_b="audiobook")
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/merge", json={"keep_id": a, "drop_id": b})
    assert resp.status_code == 400
    assert "type" in resp.get_json()["error"]

def test_dismiss_pair_guards(tmp_path):
    dbp = tmp_path / "t.db"
    a, _b = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/dismiss_pair", json={"id_a": a, "id_b": a}).status_code == 400
    assert client.post("/api/dismiss_pair", json={"id_a": a, "id_b": 999}).status_code == 400

def test_index_has_duplicates_panel():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert '<div id="dupes-panel" hidden></div>' in html

def test_edit_normalizes_genre_case(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    conn = db.connect(dbp)
    conn.execute("UPDATE enrichment SET genre=? WHERE item_id=?",
                 (db.tags_to_json(["Science Fiction"]), item_id))
    conn.commit()
    conn.close()
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post(f"/api/items/{item_id}/edit", json={"fields": {
        "genre": ["science fiction", "cozy mystery", "Cozy Mystery"],
        "authors": ["lowercase name kept"]}}).status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["genre"] == ["Science Fiction", "Cozy Mystery"]
    assert item["authors"] == ["lowercase name kept"]  # people untouched

def test_autocomplete_supports_tag_counts():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    ac = (static / "autocomplete.js").read_text(encoding="utf-8")
    # commit value must come from dataset.value, not textContent, so the
    # count span can't leak into the committed tag
    assert "dataset.value" in ac
    assert "textContent" not in ac.split("Enter")[1].split("Escape")[0]
    assert "ac-count" in ac
    assert "countsFn" in ac
    css = (static / "style.css").read_text(encoding="utf-8")
    assert ".ac-count" in css

def test_app_passes_tag_counts_to_autocomplete():
    static = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
              / "static")
    js = _viewer_js()
    assert "tagCounts" in js
    # both the tag editors and the chip filters supply a counts source
    assert js.count("tagCounts(") >= 2   # the two call sites

def test_fetch_url_falls_back_to_link_only(tmp_path, monkeypatch):
    dbp = tmp_path / "w.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    from humble_catalog import url_import
    def unavailable(conn, url):
        raise url_import.MetadataUnavailable("examplegames.com served no og:title")
    monkeypatch.setattr(url_import, "resolve", unavailable)

    resp = client.post(f"/api/items/{item_id}/fetch_url",
                       json={"url": "https://examplegames.com/p/1"})
    assert resp.status_code == 200
    cand = resp.get_json()["candidate"]
    assert cand["link_only"] is True
    assert cand["title"] == "All Systems Red"   # the item's existing name
    assert cand["url"] == "https://examplegames.com/p/1"
    assert cand["source"] == "examplegames.com"
    assert cand["authors"] is None
    assert "og:title" in cand["reason"]

def test_fetch_url_falls_back_on_network_error(tmp_path, monkeypatch):
    dbp = tmp_path / "w2.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    from humble_catalog import url_import
    def boom(conn, url):
        raise requests.ConnectionError("unreachable")
    monkeypatch.setattr(url_import, "resolve", boom)

    resp = client.post(f"/api/items/{item_id}/fetch_url",
                       json={"url": "https://examplegames.com/p/1"})
    assert resp.status_code == 200
    assert resp.get_json()["candidate"]["link_only"] is True

def test_fetch_url_rejects_bad_scheme_without_link_only(tmp_path, monkeypatch):
    dbp = tmp_path / "w3.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    from humble_catalog import url_import
    def reject(conn, url):
        raise ValueError("unsupported URL scheme 'javascript'")
    monkeypatch.setattr(url_import, "resolve", reject)

    resp = client.post(f"/api/items/{item_id}/fetch_url",
                       json={"url": "javascript:alert(1)"})
    assert resp.status_code == 400
    assert "candidate" not in resp.get_json()

def test_app_js_escapes_candidate_source():
    # source used to be a hardcoded literal from _HANDLERS; it is now a
    # hostname from a pasted URL, and urlparse does not validate netloc.
    app_js = _viewer_js()
    assert "(${c.source})" not in app_js
    assert "esc(c.source)" in app_js

def test_app_js_renders_link_only_candidates():
    app_js = _viewer_js()
    assert "link_only" in app_js
    assert "esc(c.reason)" in app_js

def test_bot_wall_403_becomes_link_only_end_to_end(tmp_path, monkeypatch):
    # The real-world case: a storefront answers the scrape with a 403 bot
    # wall. Covered in halves elsewhere (403 raises out of resolve, and a
    # RequestException degrades in the route); this pins the whole path.
    monkeypatch.setattr("time.sleep", lambda s: None)
    dbp = tmp_path / "w4.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    http = Mock()
    resp = Mock(status_code=403)
    resp.headers = {"Content-Type": "text/html"}
    resp.url = "https://examplegames.com/p/1"
    resp.raise_for_status = Mock(
        side_effect=requests.HTTPError("403 Client Error: Forbidden",
                                       response=Mock(status_code=403)))
    http.request.return_value = resp

    from humble_catalog import url_import
    real_resolve = url_import.resolve
    monkeypatch.setattr(url_import, "resolve",
                        lambda conn, url: real_resolve(conn, url, http=http))

    r = client.post(f"/api/items/{item_id}/fetch_url",
                    json={"url": "https://examplegames.com/p/1"})
    assert r.status_code == 200
    cand = r.get_json()["candidate"]
    assert cand["link_only"] is True
    assert cand["title"] == "All Systems Red"
    assert "403" in cand["reason"]
    # A bot wall is not retried: one request, no backoff.
    assert http.request.call_count == 1

def test_revert_leaves_fields_absent_from_legacy_snapshot(tmp_path):
    # pre_edit snapshots written before a field joined EDITABLE_FIELDS do
    # not contain its key. Treating that as NULL would silently wipe live
    # data on revert, so revert must skip keys it does not find.
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    # narrator is used rather than source_url because it is already in
    # EDITABLE_FIELDS: revert writes to it today, so a missing key is
    # genuinely destructive here and the test fails without the fix.
    conn = db.connect(dbp)
    conn.execute(
        "UPDATE enrichment SET series=?, narrator=?, pre_edit=? "
        "WHERE item_id=?",
        ("Edited Series", db.tags_to_json(["Sam Reader"]),
         json.dumps({"series": "Original Series"}), item_id))
    conn.commit()
    conn.close()

    assert client.post(f"/api/items/{item_id}/revert").status_code == 200

    conn = db.connect(dbp)
    row = conn.execute(
        "SELECT series, narrator FROM enrichment WHERE item_id=?",
        (item_id,)).fetchone()
    conn.close()
    assert row["series"] == "Original Series"          # present key: restored
    assert db.tags_from_json(row["narrator"]) == ["Sam Reader"]  # absent: kept

def test_edit_sets_source_url(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post(f"/api/items/{item_id}/edit",
                       json={"fields": {"source_url": "https://examplegames.com/p/1"}})
    assert resp.status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["source_url"] == "https://examplegames.com/p/1"

def test_edit_prepends_https_to_bare_source_url(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"source_url": "examplegames.com/p/1"}})
    item = client.get("/api/items").get_json()["items"][0]
    assert item["source_url"] == "https://examplegames.com/p/1"

def test_edit_rejects_javascript_source_url(tmp_path):
    # The value is rendered into an href, and esc() does not filter schemes.
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post(f"/api/items/{item_id}/edit",
                       json={"fields": {"source_url": "javascript:alert(1)"}})
    assert resp.status_code == 400
    assert "scheme" in resp.get_json()["error"].lower()
    item = client.get("/api/items").get_json()["items"][0]
    assert item["source_url"] is None      # nothing was written

def test_edit_clears_source_url(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"source_url": "https://examplegames.com/p/1"}})
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"source_url": ""}})
    item = client.get("/api/items").get_json()["items"][0]
    assert item["source_url"] is None

def test_app_js_renders_source_link_as_trailing_icon():
    # The whole title used to be the anchor, so selecting or copying a
    # title risked navigating. The link is now a trailing glyph.
    app_js = _viewer_js()
    assert "src-link" in app_js
    assert "&#x2197;" in app_js
    assert 'title="Open source page"' in app_js

def test_app_js_edits_source_url():
    app_js = _viewer_js()
    assert 'data-f="source_url"' in app_js

def test_set_user_tags(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post(f"/api/items/{item_id}/user-tags",
                       json={"tags": ["To Reread"]})
    assert resp.status_code == 200
    item = client.get("/api/items").get_json()["items"][0]
    assert item["user_tags"] == ["To Reread"]

def test_user_writes_do_not_mark_edited(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/user-tags", json={"tags": ["lent out"]})
    client.post(f"/api/items/{item_id}/comment", json={"comment": "A note."})
    item = client.get("/api/items").get_json()["items"][0]
    assert item["edited"] is False          # no snapshot: not a hand edit
    conn = db.connect(dbp)
    assert conn.execute("SELECT pre_edit FROM enrichment WHERE item_id=?",
                        (item_id,)).fetchone()["pre_edit"] is None

def test_set_user_tags_rejects_non_list(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post(f"/api/items/{item_id}/user-tags",
                       json={"tags": "nope"}).status_code == 400

def test_set_user_tags_empty_clears_to_null(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/user-tags", json={"tags": ["lent out"]})
    client.post(f"/api/items/{item_id}/user-tags", json={"tags": []})
    conn = db.connect(dbp)
    assert conn.execute("SELECT user_tags FROM items WHERE id=?",
                        (item_id,)).fetchone()["user_tags"] is None

def test_set_comment_strips_and_clears(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/comment",
                json={"comment": "  Gift from Sam.  "})
    assert client.get("/api/items").get_json()["items"][0]["user_comment"] \
        == "Gift from Sam."
    client.post(f"/api/items/{item_id}/comment", json={"comment": "   "})
    assert client.get("/api/items").get_json()["items"][0]["user_comment"] is None

def test_user_endpoints_404_on_missing_item(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/items/999/user-tags",
                       json={"tags": ["x"]}).status_code == 404
    assert client.post("/api/items/999/comment",
                       json={"comment": "x"}).status_code == 404

def test_rename_user_tag(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/user-tags", json={"tags": ["to reread"]})
    resp = client.post("/api/user-tags/rename",
                       json={"old": "to reread", "new": "reread"})
    assert resp.status_code == 200 and resp.get_json()["changed"] == 1
    assert client.get("/api/items").get_json()["items"][0]["user_tags"] \
        == ["reread"]

def test_rename_user_tag_unknown_404(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/user-tags/rename",
                       json={"old": "nope", "new": "x"}).status_code == 404

def test_rename_user_tag_blank_400(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/user-tags/rename",
                       json={"old": "", "new": "x"}).status_code == 400

def test_delete_user_tag(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/user-tags", json={"tags": ["lent out"]})
    resp = client.post("/api/user-tags/delete", json={"tag": "lent out"})
    assert resp.status_code == 200 and resp.get_json()["changed"] == 1
    assert client.get("/api/items").get_json()["items"][0]["user_tags"] == []

def test_delete_user_tag_unknown_404(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post("/api/user-tags/delete",
                       json={"tag": "nope"}).status_code == 404

def test_user_tag_rename_does_not_touch_genre(tmp_path):
    # the two vocabularies are separate pools: renaming a user tag must
    # never reach enrichment.genre, even on a same-spelled tag
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/choose", json={"candidate": 0})
    client.post(f"/api/items/{item_id}/edit", json={"fields": {"genre": ["SF"]}})
    client.post(f"/api/items/{item_id}/user-tags", json={"tags": ["SF"]})
    assert client.post("/api/user-tags/rename",
                       json={"old": "SF", "new": "space"}).get_json()["changed"] == 1
    item = client.get("/api/items").get_json()["items"][0]
    assert item["user_tags"] == ["space"]
    assert item["genre"] == ["SF"]          # genre pool untouched

def test_user_fields_never_reach_the_edit_endpoint():
    js = _viewer_js()
    # user_tags is destructured out before the /edit payload is assembled;
    # sending it would 400, since it is not in EDITABLE_FIELDS
    assert "const {user_tags, ...enrichmentTags} = editingTags;" in js
    assert "Object.assign(fields, enrichmentTags);" in js
    # the note must NOT carry .edit-field, or the selector that builds the
    # /edit payload would sweep it up
    assert 'class="edit-comment"' in js
    assert "edit-field edit-comment" not in js

def test_save_skips_edit_when_only_user_fields_changed():
    js = _viewer_js()
    # every /edit call snapshots the row and marks it hand-edited, so a
    # save that changed only the note or the user tags must not post it
    assert "function shouldPostEnrichmentEdit(item, fields)" in js
    assert "if (shouldPostEnrichmentEdit(items.find(i => i.id === +id), fields)) {" in js

def test_index_has_user_columns():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert "<th>My tags</th><th>Notes</th>" in html

def _index_html():
    return (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")

def test_user_tags_filter_is_registered():
    # one registry entry + one input is the whole contract for a new
    # filterable column; the wiring loops pick it up from there
    assert 'data-field="user_tags"' in _index_html()
    assert '<input id="f-user-tags"' in _index_html()
    assert "user_tags: {accessor: i => i.user_tags," in _viewer_js()

def test_user_tags_filter_pool_is_separate_from_genre():
    js = _viewer_js()
    # genre and user tags are independent pools: a genre "Fantasy" and a
    # personal tag "fantasy" mean different things and must stay
    # separately selectable, so no accessor may union the two
    assert "[...i.genre, ...i.user_tags]" not in js
    assert "[...i.user_tags, ...i.genre]" not in js

def test_tag_badges_tolerates_a_missing_field():
    # render() runs before loadReview/loadDupes/refreshStats, so an
    # exception here blanks the whole page, not just one column. An item
    # served without user_tags (an older API, a partial payload) must not
    # be able to do that.
    js = _viewer_js()
    assert "const tagBadges = (arr) =>\n  (arr || []).map(" in js

def test_notes_filter_is_registered():
    assert 'data-field="user_comment"' in _index_html()
    assert '<input id="f-notes"' in _index_html()
    assert "user_comment: {accessor: i => i.user_comment ? [i.user_comment] : []," in _viewer_js()

def test_notes_filter_has_no_autocomplete():
    # a dropdown suggesting whole note bodies would be useless
    assert "if (!f.textOnly) Autocomplete.attach(" in _viewer_js()

def test_chip_filter_text_is_lowercased_on_input():
    # passesChipFilters lowercases the value but not f.text, so every
    # place that writes f.text has to lowercase it first
    js = _viewer_js()
    assert "f.text = input.value.trim().toLowerCase();" in js
    assert "f.text = value.toLowerCase();" in js

def test_search_box_is_still_names_only():
    # v1.12 narrowed it deliberately; neither the notes filter nor fuzzy
    # matching may widen it. Asserted against the source rather than
    # behaviour only for the negative half -- the positive half moved to
    # test_webapp_js.py, where it runs the real visible().
    js = _viewer_js()
    assert "i.user_comment.toLowerCase().includes(q)" not in js

def test_bulk_add_and_remove(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/user-tags/bulk", json={
        "ids": [item_id], "tag": "to reread", "action": "add"})
    assert resp.status_code == 200 and resp.get_json()["changed"] == 1
    assert client.get("/api/items").get_json()["items"][0]["user_tags"] \
        == ["to reread"]
    # adding again changes nothing
    assert client.post("/api/user-tags/bulk", json={
        "ids": [item_id], "tag": "to reread",
        "action": "add"}).get_json()["changed"] == 0
    resp = client.post("/api/user-tags/bulk", json={
        "ids": [item_id], "tag": "to reread", "action": "remove"})
    assert resp.get_json()["changed"] == 1
    assert client.get("/api/items").get_json()["items"][0]["user_tags"] == []

def test_bulk_does_not_mark_rows_edited(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post("/api/user-tags/bulk", json={
        "ids": [item_id], "tag": "lent out", "action": "add"})
    assert client.get("/api/items").get_json()["items"][0]["edited"] is False

def test_bulk_validation(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    bad = [
        {"ids": [], "tag": "x", "action": "add"},              # no ids
        {"ids": "nope", "tag": "x", "action": "add"},          # ids not a list
        {"ids": [item_id], "tag": "  ", "action": "add"},      # blank tag
        {"ids": [item_id], "tag": "x", "action": "replace"},   # bad verb
        {"ids": [item_id], "tag": "x"},                        # no verb
    ]
    for payload in bad:
        assert client.post("/api/user-tags/bulk",
                           json=payload).status_code == 400, payload

def test_bulk_ignores_unknown_ids(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    # a stale page can hold ids that no longer exist; that is not an error
    resp = client.post("/api/user-tags/bulk", json={
        "ids": [item_id, 999], "tag": "to reread", "action": "add"})
    assert resp.status_code == 200 and resp.get_json()["changed"] == 1

def test_bulk_bar_is_present():
    html = _index_html()
    assert '<input id="bulk-tag"' in html
    assert '<button id="bulk-add">' in html
    assert '<button id="bulk-remove">' in html

def test_bulk_remove_is_gated_on_an_active_filter():
    # user_tags has no pre_edit snapshot, so a bulk remove cannot be undone
    assert "removeBtn.disabled = !tag || count === 0 || !filtered;" in _viewer_js()

def test_bulk_bar_lives_outside_the_table():
    # the table rebuilds via innerHTML; controls inside it would be destroyed
    html = _index_html()
    assert html.index('id="bulk-bar"') < html.index('<table id="catalog">')

def _static_dir():
    return Path(__file__).parent.parent / "humble_catalog" / "webapp" / "static"

def test_index_links_both_favicons():
    html = _index_html()
    # Browsers that understand SVG favicons take the first and ignore the
    # second; the PNG covers the rest. Absolute /static/ paths keep the
    # bare /favicon.ico request path out of it, so no Flask route is
    # needed - static_url_path="/static" already serves both.
    assert ('<link rel="icon" href="/static/favicon.svg" '
            'type="image/svg+xml">') in html
    assert ('<link rel="icon" href="/static/favicon-32.png" '
            'type="image/png" sizes="32x32">') in html
    # both files must actually ship, not just be referenced
    static = _static_dir()
    assert (static / "favicon.svg").read_text(encoding="utf-8").strip()
    assert (static / "favicon-32.png").read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"

def test_favicon_follows_the_os_theme():
    svg = (_static_dir() / "favicon.svg").read_text(encoding="utf-8")
    # The tab icon renders outside the page's CSS cascade, so it cannot
    # read style.css or the data-theme attribute the toggle sets. Its own
    # media query is the only mechanism available.
    assert "@media (prefers-color-scheme:dark)" in svg


def test_index_loads_the_fuzzy_scorer_before_the_app():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert '<script src="/static/fuzzy.js"></script>' in html
    # order matters: app.js references Fuzzy at load time
    assert html.index("fuzzy.js") < html.index("app.js")

def test_override_round_trip(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"series": "Harbor Tales"}})
    assert client.post(f"/api/items/{item_id}/override",
                       json={"override": True}).status_code == 200
    assert client.get("/api/items").get_json()["items"][0]["override"] is True
    assert client.post(f"/api/items/{item_id}/override",
                       json={"override": False}).status_code == 200
    assert client.get("/api/items").get_json()["items"][0]["override"] is False

def test_override_rejected_on_a_row_that_was_never_edited(tmp_path):
    # a machine-enriched row is already eligible for enrichment, so an
    # override on it is meaningless rather than harmless
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    assert client.post(f"/api/items/{item_id}/override",
                       json={"override": True}).status_code == 400

def test_override_rejects_a_non_boolean(tmp_path):
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"series": "Harbor Tales"}})
    assert client.post(f"/api/items/{item_id}/override",
                       json={"override": "yes"}).status_code == 400

def test_revert_of_a_re_enriched_row_hands_the_edit_back(tmp_path):
    # authorship flips: this row's snapshot holds the TYPED values, so
    # reverting makes it hand-edited again and drops any queued override
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    client.post(f"/api/items/{item_id}/edit",
                json={"fields": {"genre": ["Hand Typed"]}})
    client.post(f"/api/items/{item_id}/choose", json={"candidate": 0})
    item = client.get("/api/items").get_json()["items"][0]
    assert item["edited"] is False and item["re_enriched"] is True
    client.post(f"/api/items/{item_id}/revert")
    item = client.get("/api/items").get_json()["items"][0]
    assert item["edited"] is True and item["re_enriched"] is False
    assert item["genre"] == ["Hand Typed"] and item["override"] is False

def test_post_export_csv_narrows_the_columns(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.csv",
                       json={"ids": ids, "columns": ["type", "title"]})
    assert resp.status_code == 200
    header = resp.get_data().decode("utf-8-sig").splitlines()[0]
    assert header == "title,type"  # canonical, not requested, order

def test_post_export_drops_unknown_column_names(tmp_path):
    # Silently, unlike the CLI: a stale hc-export-columns in someone's
    # browser can outlive a COLUMNS rename by months, and a 500 there
    # punishes the user for the schema's history.
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.csv",
                       json={"ids": ids, "columns": ["title", "not_a_column"]})
    assert resp.status_code == 200
    assert resp.get_data().decode("utf-8-sig").splitlines()[0] == "title"

def test_post_export_empty_or_all_unknown_columns_means_everything(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    for cols in ([], ["nope"]):
        resp = client.post("/api/export.csv",
                           json={"ids": ids, "columns": cols})
        assert resp.status_code == 200
        header = resp.get_data().decode("utf-8-sig").splitlines()[0]
        assert header.split(",") == list(export.COLUMNS)

def test_post_export_rejects_a_non_list_columns(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.csv",
                       json={"ids": ids, "columns": "title"})
    assert resp.status_code == 400

def test_post_export_xlsx_narrows_the_columns(tmp_path):
    dbp = tmp_path / "t.db"
    ids = _seed_pair(dbp)
    client = create_app(db_path=str(dbp)).test_client()
    resp = client.post("/api/export.xlsx",
                       json={"ids": ids, "columns": ["type", "title"]})
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.get_data())).active
    assert [c.value for c in ws[1]] == ["title", "type"]

def test_header_stacks_above_the_sticky_table_head():
    # <header> is a static flex item WITH a z-index, so it establishes a
    # stacking context: every z-index inside it -- the column picker's
    # included -- is ordered only within the header, and externally the
    # whole header is one layer. The sticky <th> sits in the root context
    # (its #table-wrap ancestor has z-index auto), so if the header's
    # layer is the lower of the two, the sticky cells paint over anything
    # the header contains and no z-index on the picker can rescue it.
    # Pinned as a relationship, not as literal values, because that is the
    # part that has to stay true.
    css = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
           / "static" / "style.css").read_text(encoding="utf-8")
    def z(selector):
        block = re.search(re.escape(selector) + r"\s*\{[^}]*?z-index:\s*(\d+)",
                          css, re.S)
        assert block, f"no z-index found for {selector}"
        return int(block.group(1))
    assert z("header") > z("#catalog thead th")
    # and the autocomplete list, which hangs off <body> to escape both
    # clipping and this very stacking problem, stays above the header
    assert z(".ac-list") > z("header")

def test_api_stats_returns_every_section_with_counts(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    body = client.get("/api/stats").get_json()
    assert body["total"] == 1
    assert [s["key"] for s in body["sections"]] == [
        "type", "rating", "status", "enrichment", "gaps", "genre"]
    rows = {s["key"]: {r["label"]: r["count"] for r in s["rows"]}
            for s in body["sections"]}
    assert rows["type"]["E-books"] == 1
    assert rows["enrichment"]["Low confidence"] == 1

def test_api_stats_matches_the_report_over_the_same_db(tmp_path):
    # the route must be a reshaping of report(), never a second count
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    conn = db.connect(dbp)
    sections, total = stats.report(db.fetch_items(conn))
    conn.close()

    body = client.get("/api/stats").get_json()
    assert body["total"] == total
    assert body["sections"] == [
        {"key": key, "label": label,
         "rows": [{"label": rl, "count": n} for rl, n in rows]}
        for key, label, rows in sections]

def test_api_stats_of_an_empty_catalog_is_well_formed(tmp_path):
    dbp = tmp_path / "t.db"
    db.connect(str(dbp)).close()
    client = create_app(db_path=str(dbp)).test_client()

    body = client.get("/api/stats").get_json()
    assert body["total"] == 0
    assert len(body["sections"]) == 6

def test_index_offers_the_rating_and_enrichment_filters():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert '<select id="f-rating"' in html    # carries an aria-label too
    assert '<option value="5">★5</option>' in html
    # the enrichment section's four rows each jump to their own state, so
    # a row reading "Unmatched 4" cannot land on review's wider 8
    assert '<option value="matched">Matched</option>' in html
    assert '<option value="low_confidence">Low-confidence match</option>' in html
    assert '<option value="unmatched">Unmatched</option>' in html
    assert '<option value="pending">Pending enrichment</option>' in html
    assert '<option value="review">Needs review</option>' in html   # union stays

def test_stats_controls_are_themed_not_browser_default():
    # Same bug the form controls had: an unstyled <button> keeps the
    # browser's grey default, which glares on the dark panel. Caught in a
    # real browser -- the JS harness has no computed styles.
    css = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
           / "static" / "style.css").read_text(encoding="utf-8")
    assert ".stat-jump, .stat-show-all, .stat-edit-tags, .bundle-jump {" in css
    # must use the theme variable, never a literal, or one palette breaks
    assert "color: var(--accent);" in css


def _bundle_app(tmp_path):
    dbp = tmp_path / "catalog.db"
    conn = db.connect(dbp)
    conn.execute("INSERT INTO items (machine_name, name, type) "
                 "VALUES ('owned_examplepress', 'Unrelated Book', 'ebook')")
    conn.commit()
    conn.close()
    return create_app(db_path=str(dbp)).test_client()


_FAKE_BUNDLE = {
    "basic_data": {"human_name": "Bundle One", "currency": "EUR"},
    "tier_item_data": {"owned_examplepress": {"human_name": "Unrelated Book"},
                       "new_examplepress": {"human_name": "The Hollow Crypt"}},
    "tier_display_data": {"initial": {
        "tier_item_machine_names": ["owned_examplepress", "new_examplepress"]}},
    "tier_pricing_data": {"initial": {
        "price|money": {"currency": "EUR", "amount": 12.0}}},
}


def test_bundle_preview_route_returns_the_report(tmp_path, monkeypatch):
    from humble_catalog import bundle_preview
    monkeypatch.setattr(bundle_preview, "fetch_bundle",
                        lambda url, http=None: _FAKE_BUNDLE)
    client = _bundle_app(tmp_path)
    resp = client.post("/api/bundle-preview", json={
        "url": "https://www.humblebundle.com/books/bundle-one-books"})
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["name"] == "Bundle One"
    assert body["currency"] == "EUR"
    assert body["tiers"][0]["owned"] == 1
    assert body["tiers"][0]["new"] == 1


def test_bundle_preview_route_echoes_the_url_it_was_given(tmp_path, monkeypatch):
    from humble_catalog import bundle_preview
    monkeypatch.setattr(bundle_preview, "fetch_bundle",
                        lambda url, http=None: _FAKE_BUNDLE)
    url = "https://www.humblebundle.com/books/bundle-one-books"
    body = _bundle_app(tmp_path).post(
        "/api/bundle-preview", json={"url": url}).get_json()
    assert body["url"] == url


def test_bundle_preview_route_rejects_a_missing_url(tmp_path):
    resp = _bundle_app(tmp_path).post("/api/bundle-preview", json={})
    assert resp.status_code == 400
    assert "url required" in resp.get_json()["error"]


def test_bundle_preview_route_rejects_a_non_humble_url(tmp_path):
    resp = _bundle_app(tmp_path).post(
        "/api/bundle-preview", json={"url": "https://example.test/books/x"})
    assert resp.status_code == 400
    assert "not a HumbleBundle URL" in resp.get_json()["error"]


def test_bundle_preview_route_reports_a_dead_page_as_a_gateway_error(
        tmp_path, monkeypatch):
    from humble_catalog import bundle_preview

    def boom(url, http=None):
        raise requests.HTTPError("404 Client Error")

    monkeypatch.setattr(bundle_preview, "fetch_bundle", boom)
    resp = _bundle_app(tmp_path).post("/api/bundle-preview", json={
        "url": "https://www.humblebundle.com/books/gone"})
    assert resp.status_code == 502
    assert "404" in resp.get_json()["error"]


def test_bundle_preview_route_is_not_reachable_by_get(tmp_path):
    # A bundle URL in a query string would reach access logs and browser
    # history, and which bundles are being eyed is the same class of
    # information as which books are owned.
    resp = _bundle_app(tmp_path).get(
        "/api/bundle-preview?url=https://www.humblebundle.com/books/x")
    assert resp.status_code == 405


def test_index_has_the_bundle_preview_panel_and_input():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    assert '<div id="bundle-panel"' in html
    assert '<input id="bundle-url"' in html
    assert '<button id="bundle-go"' in html


def test_bundle_panel_scrolls_internally_like_the_other_panels():
    # Without this a 35-item adds list grows the panel without bound. It
    # can no longer push the table off screen -- they are in different
    # sections now, which is why the max-height half of this rule is gone
    # -- but a panel taller than the viewport still has to scroll itself
    # rather than the page. The JS harness has no computed styles, so this
    # is asserted against the stylesheet.
    css = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
           / "static" / "style.css").read_text(encoding="utf-8")
    assert ("#review-panel, #dupes-panel, #stats-panel, #bundle-panel "
            "{ flex: 0 1 auto;") in css
    assert "overflow: auto; }" in css


# --- Host validation (DNS rebinding defence) -------------------------------

def test_host_is_loopback_accepts_only_loopback_authorities():
    # Unit-level because the missing-Host case cannot be produced through
    # the test client, which always synthesises one.
    from humble_catalog.webapp import host_is_loopback
    assert host_is_loopback("127.0.0.1:8087")
    assert host_is_loopback("localhost:8087")
    assert host_is_loopback("localhost")
    assert host_is_loopback("[::1]:8087")
    assert host_is_loopback("[::1]")
    assert host_is_loopback("LOCALHOST:8087")      # Host is case-insensitive
    assert not host_is_loopback("evil.example.com")
    assert not host_is_loopback("evil.example.com:8087")
    assert not host_is_loopback("127.0.0.1.evil.com")
    assert not host_is_loopback("localhost.evil.com")
    assert not host_is_loopback("192.168.1.10:8087")
    assert not host_is_loopback("")
    assert not host_is_loopback(None)


def test_reads_are_refused_for_a_rebound_hostname(tmp_path):
    # The attack this closes: a page on evil.com whose DNS is re-pointed at
    # 127.0.0.1 becomes same-origin with the viewer, and same-origin policy
    # stops protecting the catalog. The Host header still says evil.com.
    dbp = tmp_path / "t.db"
    _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    resp = client.get("/api/items", headers={"Host": "evil.example.com"})
    assert resp.status_code == 403
    assert b"All Systems Red" not in resp.data


def test_writes_are_refused_before_the_body_is_even_parsed(tmp_path):
    # 403 rather than 415/404 proves the check runs ahead of routing, so a
    # foreign host cannot reach a handler by getting the content type right.
    dbp = tmp_path / "t.db"
    item_id = _seed(dbp)
    client = create_app(db_path=str(dbp)).test_client()

    resp = client.post(f"/api/items/{item_id}/rating", json={"rating": 5},
                       headers={"Host": "evil.example.com"})
    assert resp.status_code == 403


def test_loopback_callers_are_unaffected(tmp_path):
    dbp = tmp_path / "t.db"
    _seed(dbp)
    for host in ["127.0.0.1:8087", "localhost:8087", "[::1]:8087"]:
        client = create_app(db_path=str(dbp)).test_client()
        resp = client.get("/api/items", headers={"Host": host})
        assert resp.status_code == 200, host
        assert resp.get_json()["items"][0]["name"] == "All Systems Red"


def test_index_loads_the_viewer_scripts_in_dependency_order():
    from tests.js_harness import VIEWER_JS
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    # Classic scripts run in <script> order and share one global scope,
    # so order is a real dependency, not a formatting choice: app.js's
    # helpers must exist before any section defines a renderer that calls
    # them, and shell.js boots last because load() calls every renderer.
    positions = [html.index(f'/static/{p.name}"') for p in VIEWER_JS]
    assert positions == sorted(positions)
    assert VIEWER_JS[0].name == "app.js"
    assert VIEWER_JS[-1].name == "shell.js"


def test_hidden_sections_are_actually_hidden():
    css = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
           / "static" / "style.css").read_text(encoding="utf-8")
    # A section carries display:flex, which outranks the UA stylesheet's
    # [hidden] { display: none } and paints every section at once. The JS
    # stays correct throughout -- el.hidden really is true -- so no DOM
    # assertion can catch this; only the explicit guard prevents it.
    assert 'section[id^="section-"][hidden] { display: none; }' in css


def test_filter_chips_live_outside_the_collapsible_sidebar():
    html = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
            / "static" / "index.html").read_text(encoding="utf-8")
    # a collapsed sidebar must never hide a filter that is narrowing the
    # table, so the summary renders in the main column, not the aside
    aside = html[html.index('<aside id="filters"'):html.index("</aside>")]
    assert 'id="filter-chips"' not in aside
    assert 'id="filter-chips"' in html
    # search likewise stays in the header, above the tabs
    assert html.index('id="search-wrap"') < html.index('<nav id="tabs"')


def test_collapsing_the_sidebar_does_not_collapse_the_table():
    css = (Path(__file__).parent.parent / "humble_catalog" / "webapp"
           / "static" / "style.css").read_text(encoding="utf-8")
    # display:none takes the aside out of grid layout, so a collapsed
    # layout declaring two tracks auto-places the table into the FIRST.
    # With a first track of 0 the table went to zero width -- collapsing
    # the filters collapsed the catalog. One track is the fix.
    assert "#library-layout.collapsed { grid-template-columns: 1fr; }" in css
    assert "grid-template-columns: 0 1fr" not in css

def test_api_keys_matches_the_report_over_the_same_db(tmp_path):
    # A reshaping of report(), never a second count -- the /api/stats rule.
    dbp = tmp_path / "t.db"
    conn = db.connect(dbp)
    conn.execute("INSERT INTO bundles (gamekey, name, url, purchased_at) "
                 "VALUES ('kv789', 'Humble Game Bundle: Key Vault', "
                 "'https://example.invalid/kv789', '2024-01-02T00:00:00')")
    conn.execute(
        "INSERT INTO external_keys (gamekey, human_name, key_type, raw) "
        "VALUES ('kv789', 'Cinder Vale', 'steam', ?)",
        (json.dumps({"human_name": "Cinder Vale", "key_type": "steam",
                     "machine_name": "cindervale_ex"}),))
    conn.commit()
    conn.close()
    client = create_app(db_path=str(dbp)).test_client()

    body = client.get("/api/keys").get_json()
    assert body["total"] == 1
    assert body["counts"]["uncheckable"] == 1     # steam never imported
    assert [r["product"] for r in body["rows"]] == ["Cinder Vale"]

def test_api_keys_of_an_empty_catalog_is_well_formed(tmp_path):
    dbp = tmp_path / "t.db"
    db.connect(str(dbp)).close()
    client = create_app(db_path=str(dbp)).test_client()

    body = client.get("/api/keys").get_json()
    assert body["total"] == 0
    assert body["rows"] == []
    assert set(body["counts"]) == {
        "matched", "unredeemed", "uncertain", "uncheckable"}
