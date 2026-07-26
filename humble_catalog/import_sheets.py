"""Import ratings and metadata from the reference spreadsheets.

Read-only on the spreadsheets; strictly gap-filling on the catalog
(see docs/superpowers/specs/2026-07-18-ratings-import-design.md)."""
import difflib
import re
from pathlib import Path
from openpyxl import load_workbook
from humble_catalog import db
from humble_catalog.progress import Progress

DEFAULT_FILES = (
    "Reference spreadsheets/Audiobooks (loose collection).xlsx",
    "Reference spreadsheets/E-books (loose collection).xlsx",
)

# Sheet header (lowercased, colon stripped) -> canonical field. Fields
# mapping to None are recognized but dropped: bundles and publishers are
# authoritative from Humble orders. 'setting' is one sheet's
# genre-equivalent column.
HEADERS = {
    "name": "title", "genre": "genre", "setting": "genre",
    "series": "series", "number in the series": "series_number",
    "author": "authors", "narrator": "narrator", "narator": "narrator",
    "rating": "rating", "bundle": None, "humble": None, "publisher": None,
}

def _canon(cell):
    if not isinstance(cell, str):
        return None
    return HEADERS.get(cell.strip().rstrip(":").strip().lower(), None)

def _cell(value):
    """Trimmed cell value, or None when effectively empty. "N/A" counts
    as empty (the sheets use it for "no genre"/"no number"). Note that a
    numeric 0 survives (0 == neither None nor ""): a 0 rating means
    "unrated" and is skipped later with that meaning, not lost here."""
    if isinstance(value, str):
        value = value.strip()
        if value.lower() == "n/a":
            return None
    return value if value not in (None, "") else None

def read_workbook(path):
    """-> (rows, skipped_sheet_names). Row dicts carry sheet, title,
    rating (may be None) and any non-empty canonical fields."""
    wb = load_workbook(path, data_only=True, read_only=True)
    rows, skipped = [], []
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        fields = [_canon(c) for c in header] if header else []
        if "title" not in fields:
            skipped.append(ws.title)
            continue
        for raw in ws.iter_rows(min_row=2, values_only=True):
            row = {"sheet": ws.title, "rating": None}
            for field, value in zip(fields, raw):
                if field is None:
                    continue
                value = _cell(value)
                if value is not None:
                    row[field] = value
            title = row.get("title")
            if title is None or (isinstance(title, str) and not title):
                continue
            rows.append(row)
    wb.close()
    return rows, skipped

_PUNCT = re.compile(r"[^\w\s]")

def norm_title(value):
    """Match key for a title: str() (floats like 1632.0 -> '1632'),
    lowercase, punctuation stripped, whitespace collapsed."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _PUNCT.sub(" ", str(value).lower())
    return " ".join(text.split())

def type_scope(path):
    """Audiobooks workbook matches audiobook items; anything else
    matches ebooks and comics."""
    if "audiobook" in Path(path).name.lower():
        return ("audiobook",)
    return ("ebook", "comic")

def build_index(conn, types):
    """-> (index, display) within the type scope: normalized catalog
    name -> [item_id, ...], and normalized name -> one original name
    (for human-readable report suggestions)."""
    index, display = {}, {}
    qmarks = ",".join("?" * len(types))
    for r in conn.execute(f"SELECT id, name FROM items WHERE type IN ({qmarks})", types):
        key = norm_title(r["name"])
        index.setdefault(key, []).append(r["id"])
        display.setdefault(key, r["name"])
    return index, display

def match(index, display, title):
    """-> (status, item_id, suggestion): ('matched', id, None) on a
    unique hit; ('ambiguous', None, None) on several; ('unmatched',
    None, closest-catalog-name-or-None) on zero. Suggestions are
    display-only, never applied."""
    ids = index.get(norm_title(title), [])
    if len(ids) == 1:
        return "matched", ids[0], None
    if ids:
        return "ambiguous", None, None
    close = difflib.get_close_matches(norm_title(title), index.keys(), n=1)
    return "unmatched", None, (display[close[0]] if close else None)

_TAG_FIELDS = ("genre", "authors", "narrator")     # sheet fields that are tags
_SCALAR_FIELDS = ("series", "series_number")

def import_row(conn, item_id, row):
    """Gap-fill one matched row. Returns (ratings_set, fields_filled).
    Only empty catalog fields are written; metadata goes through
    db.apply_hand_edit so the row gets the edited badge and a revert
    snapshot — but only when something is actually written."""
    ratings_set = 0
    rating = row.get("rating")
    if isinstance(rating, (int, float)) and rating > 0:
        cur = conn.execute("SELECT my_rating FROM items WHERE id=?",
                           (item_id,)).fetchone()
        if cur["my_rating"] is None:
            conn.execute("UPDATE items SET my_rating=? WHERE id=?",
                         (int(rating), item_id))
            conn.commit()
            ratings_set = 1
    current = conn.execute(
        "SELECT genre, authors, narrator, series, series_number "
        "FROM enrichment WHERE item_id=?", (item_id,)).fetchone()
    fields = {}
    for f in _TAG_FIELDS:
        if row.get(f) is not None and not db.tags_from_json(current[f]):
            vals = [str(row[f])]
            if f == "genre":
                vals = db.normalize_tags(conn, db.GENRE, vals)
            fields[f] = db.tags_to_json(vals)
    for f in _SCALAR_FIELDS:
        if row.get(f) is not None and current[f] is None:
            value = row[f]
            if f == "series_number":
                try:
                    fields[f] = float(value)
                except (TypeError, ValueError):
                    continue    # not a number ("N/A" etc.) -> skip the field
            else:
                fields[f] = str(value)
    if fields:
        db.apply_hand_edit(conn, item_id, fields)
    return ratings_set, len(fields)

def run(paths=None, db_path="catalog.db"):
    """Import every spreadsheet, reporting live progress in two phases.

    Reading is separated from importing so the row phase can show a real
    total (and therefore an ETA) - openpyxl only knows how many rows a
    workbook holds once it has parsed it. That costs holding every
    workbook's rows at once rather than one at a time, which is why it is
    worth saying out loud; DEFAULT_FILES is a handful of files.
    """
    paths = list(paths) if paths else [p for p in DEFAULT_FILES if Path(p).exists()]
    conn = db.connect(db_path)
    totals = {"rows": 0, "matched": 0, "ratings": 0, "filled": 0, "complete": 0}
    unmatched, ambiguous, skipped_sheets = [], [], []
    try:
        sheets = []
        reading = Progress(conn, "import-sheets", total=len(paths),
                           phase="Workbook", echo=False)
        for path in paths:
            reading.step(Path(path).name)  # named BEFORE the slow parse
            rows, skipped = read_workbook(path)
            skipped_sheets += [f"{Path(path).name}/{s}" for s in skipped]
            sheets.append((path, rows))
        reading.detach()
        prog = Progress(conn, "import-sheets", phase="Row", echo=False,
                        total=sum(len(rows) for _, rows in sheets),
                        tallies=("matched", "ratings", "filled",
                                 "unmatched", "ambiguous"))
        for path, rows in sheets:
            index, display = build_index(conn, type_scope(path))
            for row in rows:
                prog.step(row["title"])
                totals["rows"] += 1
                status, item_id, suggestion = match(index, display, row["title"])
                where = f"[{Path(path).name}/{row['sheet']}] {row['title']}"
                if status == "ambiguous":
                    ambiguous.append(where)
                    prog.count("ambiguous")
                elif status == "unmatched":
                    unmatched.append(
                        where + (f" — closest: {suggestion}" if suggestion else ""))
                    prog.count("unmatched")
                else:
                    totals["matched"] += 1
                    prog.count("matched")
                    ratings, filled = import_row(conn, item_id, row)
                    totals["ratings"] += ratings
                    totals["filled"] += filled
                    prog.count("ratings", ratings)
                    prog.count("filled", filled)
                    if ratings == 0 and filled == 0:
                        totals["complete"] += 1
        prog.finish(f"{totals['rows']} rows read, {totals['matched']} matched, "
                    f"{totals['ratings']} ratings set, {totals['filled']} fields "
                    f"filled, {totals['complete']} already complete")
    finally:
        conn.close()
    for label, entries in (("Skipped sheets (no Name: column)", skipped_sheets),
                           ("Ambiguous (several catalog matches)", ambiguous),
                           ("Unmatched", unmatched)):
        if entries:
            print(f"\n{label}:")
            for e in entries:
                print(f"  {e}")
