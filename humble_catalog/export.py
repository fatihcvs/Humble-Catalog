import csv
import re
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from humble_catalog import db

COLUMNS = ("title", "type", "publisher", "authors", "genre", "series",
           "series_number", "narrator", "illustrator", "my_rating",
           "external_rating", "rating_source", "formats", "bundles",
           "first_purchased", "status", "edited", "user_tags", "user_comment",
           "read_status")

# read_status is stored as a stable snake_case key; the export shows the
# human label, matching how the rest of the export favours readable cells.
READ_STATUS_LABELS = {"want_to_read": "Want to read", "unread": "Unread",
                      "reading": "Reading", "read": "Read", "dnf": "DNF"}

# A long user_comment or bundles join would otherwise stretch one column
# off the screen; there is no autofit in the format, the width is a
# number someone has to choose.
WIDTH_CAP = 60

def _columns(requested):
    # The column policy, in one place because three callers depend on it
    # -- the two routes and the CLI -- and a second copy is how they
    # start disagreeing about what a valid selection is.
    #
    # The caller does NOT own column order, which is the opposite of the
    # rule for rows in _select. That is not an inconsistency: on-screen
    # row order is something the user built with sorting and relevance
    # ranking, so it is information. Column order is not -- nobody drags
    # the checkboxes -- so forcing canonical order costs nothing and buys
    # validation and de-duplication in the same comprehension.
    #
    # Empty means everything, so `columns=None` keeps every existing
    # caller working by construction, the same trick `ids=None` pulls.
    if not requested:
        return COLUMNS
    return tuple(c for c in COLUMNS if c in requested) or COLUMNS

def _select(conn, ids):
    # The rows to write, in the order they must be written.
    #
    # ids=None is the whole catalog in fetch_items' title order. With ids
    # the CALLER owns the order -- the viewer posts rows in the order it
    # shows them, relevance ordering included -- so this must never
    # re-sort. Unknown ids are skipped rather than fatal: the export is
    # read-only, and a row that vanished since the page loaded should not
    # cost the whole download. Callers return len() of this list, so the
    # reported count stays truthful either way.
    #
    # Both writers go through here. Copying the policy into the second
    # one is how two exports start disagreeing about what a stale id means.
    items = db.fetch_items(conn)
    if ids is None:
        return items
    by_id = {i["id"]: i for i in items}
    return [by_id[i] for i in ids if i in by_id]

def _row(item, columns):
    purchased = [b["purchased_at"] for b in item["bundles"] if b["purchased_at"]]
    values = dict(item,
                  title=item["name"],
                  genre="; ".join(item["genre"]),
                  authors="; ".join(item["authors"]),
                  narrator="; ".join(item["narrator"]),
                  illustrator="; ".join(item["illustrator"]),
                  user_tags="; ".join(item["user_tags"]),
                  formats="; ".join(item["formats"]),
                  bundles="; ".join(b["name"] for b in item["bundles"]),
                  first_purchased=(date.fromisoformat(min(purchased)[:10])
                                   if purchased else None),
                  edited="yes" if item["edited"] else None,
                  read_status=READ_STATUS_LABELS.get(
                      item["read_status"], item["read_status"]))
    # The values dict above is built the same way whatever the projection:
    # branching per column would save nothing measurable.
    return ["" if values[c] is None else values[c] for c in columns]

def write_csv(conn, fh, ids=None, columns=None):
    # Caller owns encoding: open files with encoding="utf-8-sig", newline="".
    # Row selection and ordering policy live in _select.
    #
    # first_purchased arrives as a datetime.date, not a string: csv.writer
    # calls str() on non-string values and str(date(2019, 3, 2)) is exactly
    # "2019-03-02" -- the same bytes this wrote before write_xlsx needed a
    # typed cell. Do not "fix" that by formatting in _row; the byte
    # identity is pinned by test_first_purchased_is_a_date_that_csv_stringifies.
    cols = _columns(columns)
    writer = csv.writer(fh)
    writer.writerow(cols)
    items = _select(conn, ids)
    for item in items:
        writer.writerow(_row(item, cols))
    return len(items)

# The set openpyxl refuses, matching its own ILLEGAL_CHARACTERS_RE.
# Tab (\x09), newline (\x0a) and carriage return (\x0d) are legal in a
# cell and are deliberately absent: stripping them would silently
# reformat a multi-line note.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def _clean(value):
    # Only str values can carry them; dates and numbers pass through.
    return _ILLEGAL.sub("", value) if isinstance(value, str) else value

def _style(ws, columns):
    # Everything that makes the workbook a working surface rather than a
    # data dump. Deliberately minimal: no colour, no conditional
    # formatting, no wrapped text -- decoration the person actually using
    # the sheet is better placed to choose.
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
    # Looked up in the columns ACTUALLY being written, not in COLUMNS: a
    # positional assumption against the global stopped being true the
    # moment the column set became a runtime choice. Absent means there
    # is simply no date column to format.
    if "first_purchased" in columns:
        date_col = columns.index("first_purchased") + 1
        for row in ws.iter_rows(min_row=2, min_col=date_col, max_col=date_col):
            # The cell holds a real date; this is only how Excel renders it.
            row[0].number_format = "YYYY-MM-DD"
    for idx, column in enumerate(ws.iter_cols(), start=1):
        # Row 1 is in `column`, so the header's own width is the floor
        # for free -- an empty column stays as wide as its name.
        widest = max(len(str(c.value)) for c in column if c.value is not None)
        ws.column_dimensions[get_column_letter(idx)].width = min(
            widest + 2, WIDTH_CAP)

def write_xlsx(conn, fh, ids=None, columns=None):
    # Same ids contract and same return value as write_csv, but `fh` must
    # be a BINARY handle: openpyxl owns the file's encoding, so there is
    # no utf-8-sig here and no text-mode equivalent. Every caller has to
    # know this, which is why it is the first thing said.
    cols = _columns(columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog"
    ws.append(list(cols))
    items = _select(conn, ids)
    for item in items:
        # Silent: a download that fails for a reason the user can neither
        # see nor fix is worse than one missing an unprintable character.
        ws.append([_clean(v) for v in _row(item, cols)])
    _style(ws, cols)
    wb.save(fh)
    return len(items)
