"""Known-answer battery for the webapp-export-routes inventory row.

Covers `/api/export.csv` and `/api/export.xlsx` in both their GET and
POST forms, and the `_export_request` body handling they share.

The two routes exist to hand the owner their own catalog, so the cases
assert CONTENT rather than status codes: which rows come back, in which
order, under which column set. A liveness probe that checked only for a
200 and a Content-Disposition header would certify an export that
silently returned the whole catalog when the viewer asked for three rows.

Every title is invented, from docs/TEST-DATA.md. Fresh database per case.
"""
import csv
import io
import pathlib
import sys
import tempfile

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))

from humble_catalog import db  # noqa: E402
from humble_catalog.webapp import create_app  # noqa: E402

PASS, FAIL = [], []


def check(label, got, want):
    (PASS if got == want else FAIL).append(label)
    if got != want:
        print(f"  FAIL {label}\n       got  {got!r}\n       want {want!r}")


TITLES = ["Salt and Sextant", "Nightjar Post", "The Quiet Harbor: A Novel"]


def seeded():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    path = pathlib.Path(tmp.name)
    conn = db.connect(path)
    ids = []
    for i, title in enumerate(TITLES):
        cur = conn.execute(
            "INSERT INTO items (machine_name, name, type) VALUES (?,?,?)",
            (f"mn_{i}", title, "ebook"))
        ids.append(cur.lastrowid)
        conn.execute("INSERT INTO enrichment (item_id) VALUES (?)",
                     (cur.lastrowid,))
    conn.commit()
    conn.close()
    return create_app(db_path=str(path)).test_client(), ids


def csv_rows(resp):
    text = resp.data.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def xlsx_rows(resp):
    book = load_workbook(io.BytesIO(resp.data))
    return [[c.value for c in row] for row in book.active.iter_rows()]


def title_column(rows):
    # Column names are lowercase in this export; looking up the index
    # rather than assuming one, and failing loudly if it is absent, so a
    # renamed column cannot make these cases read a neighbouring field.
    header = rows[0]
    assert "title" in header, header
    idx = header.index("title")
    return [r[idx] for r in rows[1:]]


# ---------------------------------------------------------------- GET form

def case_get_returns_the_whole_catalog():
    client, _ids = seeded()
    rows = csv_rows(client.get("/api/export.csv"))
    check("GET csv returns every row", sorted(title_column(rows)),
          sorted(TITLES))


def case_get_xlsx_returns_the_whole_catalog():
    client, _ids = seeded()
    rows = xlsx_rows(client.get("/api/export.xlsx"))
    check("GET xlsx returns every row", sorted(title_column(rows)),
          sorted(TITLES))


def case_both_formats_agree_on_the_rows():
    # The invariant that matters: the two routes reshape one report and
    # must not be able to disagree about which rows it holds.
    client, _ids = seeded()
    from_csv = sorted(title_column(csv_rows(client.get("/api/export.csv"))))
    from_xlsx = sorted(title_column(xlsx_rows(client.get("/api/export.xlsx"))))
    check("csv and xlsx return the same rows", from_csv, from_xlsx)


def case_the_csv_carries_the_bom():
    # utf-8-sig, the same bytes the CLI export writes, so a spreadsheet
    # opening it does not mangle accented titles.
    client, _ids = seeded()
    check("the csv begins with the utf-8 BOM",
          client.get("/api/export.csv").data[:3], b"\xef\xbb\xbf")


def case_the_attachment_names_are_stable():
    client, _ids = seeded()
    for route, name in [("/api/export.csv", "catalog.csv"),
                        ("/api/export.xlsx", "catalog.xlsx")]:
        resp = client.get(route)
        check(f"{route} is served as an attachment named {name}",
              resp.headers["Content-Disposition"],
              f"attachment; filename={name}")


# --------------------------------------------------------------- POST form

def case_post_returns_only_the_ids_it_was_given():
    client, ids = seeded()
    resp = client.post("/api/export.csv", json={"ids": [ids[1]]})
    check("POST with one id returns exactly that row",
          title_column(csv_rows(resp)), ["Nightjar Post"])


def case_post_keeps_the_order_the_viewer_sent():
    # The filter predicate lives in the browser, so the row ORDER has to
    # arrive from there; re-sorting here would silently discard the sort
    # the owner is looking at.
    client, ids = seeded()
    reversed_ids = list(reversed(ids))
    rows = csv_rows(client.post("/api/export.csv", json={"ids": reversed_ids}))
    check("POST returns rows in the order the ids were given",
          title_column(rows), list(reversed(TITLES)))


def case_post_with_an_empty_id_list_returns_no_rows():
    # Distinct from GET's whole catalog: an empty selection is a real
    # answer, not a missing one.
    client, _ids = seeded()
    rows = csv_rows(client.post("/api/export.csv", json={"ids": []}))
    check("an empty id list exports a header and no rows",
          len(rows[1:]), 0)


def case_post_columns_selects_the_column_set():
    client, ids = seeded()
    rows = csv_rows(client.post("/api/export.csv",
                                json={"ids": ids, "columns": ["title"]}))
    check("a one-column request returns one column", rows[0], ["title"])
    check("and still returns every requested row", len(rows[1:]), len(ids))


def case_post_columns_at_two_values_changes_the_output():
    # The documented parameter, exercised at two values that must differ.
    client, ids = seeded()
    one = csv_rows(client.post("/api/export.csv",
                               json={"ids": ids, "columns": ["title"]}))[0]
    two = csv_rows(client.post("/api/export.csv",
                               json={"ids": ids,
                                     "columns": ["title", "type"]}))[0]
    check("one column versus two changes the header", (one, two),
          (["title"], ["title", "type"]))


def case_an_unknown_column_name_is_dropped_not_refused():
    # Documented asymmetry: an unknown NAME is dropped because stored
    # browser state outlives a schema rename, while a non-list is a 400.
    client, ids = seeded()
    resp = client.post("/api/export.csv",
                       json={"ids": ids, "columns": ["title", "no_such_column"]})
    check("an unknown column name is not an error", resp.status_code, 200)
    check("and is simply absent from the header",
          csv_rows(resp)[0], ["title"])


def case_an_unknown_id_is_dropped_not_refused():
    client, ids = seeded()
    rows = csv_rows(client.post("/api/export.csv",
                                json={"ids": [ids[0], 9999]}))
    check("an unknown id contributes no row",
          title_column(rows), ["Salt and Sextant"])


def case_a_malformed_body_is_refused_by_both_formats():
    client, _ids = seeded()
    bad = [("ids is not a list", {"ids": "1,2,3"}),
           ("ids is a number", {"ids": 5}),
           ("columns is not a list", {"ids": [], "columns": "title"})]
    for route in ["/api/export.csv", "/api/export.xlsx"]:
        for label, body in bad:
            resp = client.post(route, json=body)
            check(f"{route} refuses {label}", resp.status_code, 400)
            check(f"{route} explains {label}",
                  resp.get_json() is not None, True)


def case_both_formats_refuse_identically():
    # _export_request is shared so the two cannot drift on what a
    # malformed body means; asserted rather than assumed.
    client, _ids = seeded()
    body = {"ids": "not a list"}
    a = client.post("/api/export.csv", json=body)
    b = client.post("/api/export.xlsx", json=body)
    check("both formats give the same refusal",
          (a.status_code, a.get_json()), (b.status_code, b.get_json()))


def case_an_absent_ids_key_is_refused():
    client, _ids = seeded()
    check("a POST with no ids key is refused",
          client.post("/api/export.csv", json={}).status_code, 400)


def case_an_empty_catalog_exports_a_header():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    path = pathlib.Path(tmp.name)
    db.connect(path).close()
    client = create_app(db_path=str(path)).test_client()
    rows = csv_rows(client.get("/api/export.csv"))
    check("an empty catalog still exports its header row",
          (len(rows) >= 1, len(rows[1:])), (True, 0))


CASES = [v for k, v in sorted(globals().items()) if k.startswith("case_")]

if __name__ == "__main__":
    for fn in CASES:
        try:
            fn()
        except Exception as exc:                            # noqa: BLE001
            FAIL.append(fn.__name__)
            print(f"  FAIL {fn.__name__} raised: "
                  f"{type(exc).__name__}: {exc}")
    total = len(PASS) + len(FAIL)
    print(f"webapp-export-routes: {len(PASS)}/{total}")
    sys.exit(1 if FAIL else 0)
